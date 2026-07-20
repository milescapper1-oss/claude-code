#!/usr/bin/env python3
"""
Saunders valve-BD contact pipeline — single source of truth is data/contacts_source.csv.

Add new contacts (e.g. a LinkedIn Sales Navigator export) by APPENDING rows to
data/contacts_source.csv, then run:  python scripts/build.py

It regenerates, in data/:
  contacts_research.json      full structured records (name..tier..persona..layer)
  contacts_ranked.csv         Name | Title | Company | Tier | Site/Region | Rationale
  contacts_persona_match.csv  every contact -> specific persona + source + where-to-meet
  contacts_all.xlsx           multi-tab workbook (Summary | Decision-Makers | per-account)

No hardcoded paths — resolves everything relative to this file, so it works in any
clone / VS Code workspace. Only dependency: openpyxl  (pip install -r requirements.txt)
"""
import csv, json, sys
from pathlib import Path
from collections import Counter

DATA = Path(__file__).resolve().parent.parent / "data"
SRC = DATA / "contacts_source.csv"

# ---------------------------------------------------------------- classification
def classify(function, title):
    """Return (tier, rationale). Tier 1 primary / 2 influencer / 3 exec / 0 de-prioritise."""
    f = (function or "").lower(); t = (title or "").lower(); fb = f + " | " + t
    if "unknown" in f:
        return 0, "Role/title unknown — verify function before ranking."
    if ("medical science liaison" in fb or "medical affairs" in fb or "country gm" in f
            or "regulatory affairs" in fb or ("research" in fb and "manufactur" not in fb)):
        return 0, "Medical/commercial/research — low relevance to valve decisions (de-prioritise)."
    if any(k in fb for k in ("utilities", "technical services", "cip", "sip", "wfi", "clean util")):
        return 1, "Utilities / CIP-SIP — owns clean-utility systems where diaphragm valves spec in."
    if any(k in fb for k in ("capital project", "construction", "expansion", "project delivery", "capex")):
        return 1, "Capital-project delivery — owns valve/equipment spec & procurement on an active build."
    if "msat" in fb or "technical service" in fb:
        return 1, "MSAT / tech ops — process ownership feeding valve & equipment selection."
    if any(k in fb for k in ("reliability", "maintenance", "asset")):
        return 1, "Asset reliability & maintenance — the predictive-maintenance / lifecycle buyer."
    if "engineering" in fb or "commissioning" in fb or "c&q" in fb or "cqv" in fb:
        return 1, "Engineering / C&Q — specifies and standardises process valves & equipment."
    if any(k in fb for k in ("site lead", "site head", "site director", "plant manager", "campus lead", "general manager", "gm ")):
        return 1, "Site/plant leadership — owns local asset & standardisation decisions."
    if any(k in fb for k in ("quality", "qa", "validation", "qualif", "compliance")):
        return 2, "QA / validation — influences equipment qualification and change control."
    if any(k in fb for k in ("procurement", "sourcing", "supply chain", "supply", "category")):
        return 2, "Sourcing / procurement — commercial gatekeeper for valve & equipment contracts."
    if any(k in fb for k in ("process development", "process dev", "advanced manufactur", "digital", "automation", "technology")):
        return 2, "Process development / advanced-mfg — specs process equipment early."
    if "executive" in f or any(k in t for k in ("evp", "svp", "president", "chief", "head of")):
        return 3, "Executive manufacturing/operations sponsor — budget, direction, executive access."
    return 2, "Manufacturing-adjacent influencer — moderate valve/equipment relevance."


def level(title):
    """Inferred authority level (proxy for seniority — NOT a real reporting line)."""
    t = " " + (title or "").lower() + " "
    if any(k in t for k in (" evp", " svp", " vp ", "vice president", "president", "chief", " ceo ")):
        return (6, "Executive (VP+)")
    if any(k in t for k in ("head of", "campus lead", "site head", "site director", "site lead", "general manager", " gm ")):
        return (5, "Site / Function Head")
    if "associate director" in t:
        return (3, "Assoc Director")
    if "director" in t:
        return (4, "Director")
    if any(k in t for k in ("senior manager", "sr manager", "principal engineer", "sr principal")):
        return (3, "Sr Manager / Principal")
    if "manager" in t or " lead " in t or "team lead" in t:
        return (3, "Manager / Lead")
    if any(k in t for k in ("senior", "sr ", "sr.", "specialist", "advisor", " sme", "principal")):
        return (2, "Senior IC")
    return (1, "Engineer / IC")


def layer(title, region):
    lv = level(title)[0]
    if lv >= 6:
        return "1. Leadership"
    if lv >= 4 or (region or "").lower() == "global" or "global" in (title or "").lower():
        return "2. Global / Functional"
    return "3. Site / Working"


PERSONAS = [
    ("T1 · Utilities / CIP-SIP", 1, ["utilit", "cip", "sip", "wfi", "clean util", "black util", "steam", "water"]),
    ("T1 · Capital Projects / Construction", 1, ["capital project", "capital steward", "construction", "expansion", "project delivery", "major project", "capex"]),
    ("T1 · Engineering Standards", 1, ["engineering standard", "standards manager", "design authority", "technical governance"]),
    ("T1 · Asset Reliability & Maintenance", 1, ["reliab", "maintenance", "asset lifecycle", "asset management", "cmms"]),
    ("T1 · MSAT / Technical Service", 1, ["msat", "technical service", "tech transfer", "technology transfer"]),
    ("T1 · Commissioning & Qualification (C&Q)", 1, ["c&q", "cqv", "commissioning", "qualification"]),
    ("T1 · Site Leadership", 1, ["site head", "site lead", "site director", "plant manager", "general manager", "campus lead"]),
    ("T1 · Engineering (Project/Process)", 1, ["engineer", "engineering"]),
    ("T3 · Executive (VP/SVP/Head)", 3, ["evp", "svp", "senior vice", "vice president", "president", "chief", "group vp", "global head", "head of manufactur", "head of supply"]),
    ("T2 · QA & Validation", 2, ["quality", "qa ", "validation", "compliance", "gmp"]),
    ("T2 · Process Development / Advanced Mfg", 2, ["process development", "process dev", "advanced manuf", "digital", "automation", "pharma 4"]),
    ("T2 · Strategic Sourcing / Procurement", 2, ["sourcing", "procurement", "category", "supply chain", "purchasing"]),
    ("T2 · External Mfg / CMO", 2, ["external manuf", "cmo", "contract manuf", "external supply"]),
    ("De-prioritise (commercial/clinical/etc.)", 0, ["commercial", "market access", "marketing", "clinical", "medical affairs", "regulatory affairs", "investor", "sales", "finance"]),
]

def persona(title):
    t = " " + (title or "").lower() + " "
    for name, tier, kws in PERSONAS:
        if any(k in t for k in kws):
            return name, tier
    return "Unclassified", 1


# ---------------------------------------------------------------- load + enrich
def load():
    if not SRC.exists():
        sys.exit(f"missing {SRC} — see README.md")
    rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
    for r in rows:
        r["tier"], r["rationale"] = classify(r.get("function", ""), r.get("title", ""))
        r["level"], r["level_label"] = level(r.get("title", ""))
        r["layer"] = layer(r.get("title", ""), r.get("region", ""))
        r["persona"], r["persona_tier"] = persona(r.get("title", ""))
        site, reg = r.get("site", ""), r.get("region", "")
        r["sr"] = f"{site} ({reg})" if reg and reg.lower() not in site.lower() else site
    return rows


TIER_LABEL = {1: "Tier 1", 2: "Tier 2", 3: "Tier 3", 0: "De-prioritise"}

def main():
    rows = load()
    # ---- contacts_research.json
    json.dump(rows, open(DATA / "contacts_research.json", "w"), indent=1, ensure_ascii=False)
    # ---- contacts_ranked.csv
    trank = lambda t: 99 if t == 0 else t
    with open(DATA / "contacts_ranked.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["Name", "Title", "Company", "Tier", "Site/Region", "Rationale"])
        for r in sorted(rows, key=lambda r: (trank(r["tier"]), r["company"], r["name"])):
            w.writerow([r["name"], r["title"], r["company"], TIER_LABEL[r["tier"]], r["sr"], r["rationale"]])
    # ---- contacts_persona_match.csv
    with open(DATA / "contacts_persona_match.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["Name", "Title", "Company", "Persona", "Persona Tier", "Region", "Source", "Where to Meet"])
        for r in sorted(rows, key=lambda r: (trank(r["persona_tier"]), r["company"], r["persona"])):
            w.writerow([r["name"], r["title"], r["company"], r["persona"],
                        f"Tier {r['persona_tier']}" if r["persona_tier"] else "De-prioritise",
                        r["region"], r["source"], r["meet_at"]])
    build_xlsx(rows)
    # ---- console summary
    print(f"contacts: {len(rows)}")
    print("by account:", dict(Counter(r["company"] for r in rows)))
    print("by tier:", {TIER_LABEL[k]: v for k, v in sorted(Counter(r["tier"] for r in rows).items())})
    print("decision-makers (Director+):", len([r for r in rows if r["level"] >= 4]))
    print("source:", dict(Counter(r["source"] for r in rows)))


def build_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("openpyxl not installed — skipping .xlsx (pip install -r requirements.txt)"); return
    A = "Arial"
    hdr_fill = PatternFill("solid", fgColor="1F3864"); hdr_font = Font(name=A, bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    tier_fill = {"Tier 1": PatternFill("solid", fgColor="C6EFCE"), "Tier 2": PatternFill("solid", fgColor="DDEBF7"),
                 "Tier 3": PatternFill("solid", fgColor="FFF2CC"), "De-prioritise": PatternFill("solid", fgColor="F2F2F2")}
    LAYER_FILL = {"1. Leadership": PatternFill("solid", fgColor="F4B183"),
                  "2. Global / Functional": PatternFill("solid", fgColor="FFE699"),
                  "3. Site / Working": PatternFill("solid", fgColor="E2EFDA")}
    LEVEL_FILL = {6: PatternFill("solid", fgColor="F4B183"), 5: PatternFill("solid", fgColor="F8CBAD"),
                  4: PatternFill("solid", fgColor="FFE699"), 3: PatternFill("solid", fgColor="FFF2CC")}
    HEADERS = ["Name", "Title", "Company", "Org Layer", "Role Level", "Tier", "Site / Region", "Function", "Source", "Where to Meet", "Rationale (valve-BD)"]
    WIDTHS = [22, 44, 18, 22, 22, 12, 28, 22, 14, 24, 44]

    def tbl(ws, subset):
        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(1, c, h); cell.fill = hdr_fill; cell.font = hdr_font; cell.border = border
        for i, r in enumerate(subset, 2):
            tl = TIER_LABEL[r["tier"]]
            vals = [r["name"], r["title"], r["company"], r["layer"], r["level_label"], tl, r["sr"],
                    r["function"], r["source"], r["meet_at"], r["rationale"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(i, c, v); cell.font = Font(name=A, size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 11))); cell.border = border
                if c == 4: cell.fill = LAYER_FILL[r["layer"]]
                if c == 5 and r["level"] in LEVEL_FILL: cell.fill = LEVEL_FILL[r["level"]]
                if c == 6: cell.fill = tier_fill[tl]
        ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:K{len(subset)+1}"
        for c, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[chr(64 + c)].width = w

    wb = Workbook()
    accounts = sorted({r["company"] for r in rows})
    trank = lambda t: 99 if t == 0 else t

    sm = wb.active; sm.title = "Summary"
    sm["A1"] = "Saunders Diaphragm-Valve BD — Target-Account Contact Base"; sm["A1"].font = Font(name=A, bold=True, size=13)
    sm["A2"] = f"{len(rows)} named contacts. Tiers = influence over valve spec/procurement/lifecycle. Regenerate with scripts/build.py."
    sm["A2"].font = Font(name=A, size=10, italic=True, color="595959")
    acc_ct = Counter(r["company"] for r in rows); tier_ct = Counter(TIER_LABEL[r["tier"]] for r in rows); src_ct = Counter(r["source"] for r in rows)
    sm["A4"] = "By Account"; sm["A4"].font = Font(name=A, bold=True, size=12)
    for c, h in enumerate(["Account", "Contacts"], 1):
        cc = sm.cell(5, c, h); cc.font = Font(name=A, bold=True, color="FFFFFF"); cc.fill = hdr_fill
    for i, acc in enumerate(accounts, 6):
        sm.cell(i, 1, acc); sm.cell(i, 2, acc_ct.get(acc, 0))
    tr = 6 + len(accounts); sm.cell(tr, 1, "TOTAL").font = Font(name=A, bold=True); sm.cell(tr, 2, len(rows)).font = Font(name=A, bold=True)
    sm["D4"] = "By Tier"; sm["D4"].font = Font(name=A, bold=True, size=12)
    for c, h in zip((4, 5), ["Tier", "Contacts"]):
        cc = sm.cell(5, c, h); cc.font = Font(name=A, bold=True, color="FFFFFF"); cc.fill = hdr_fill
    for i, t in enumerate(["Tier 1", "Tier 2", "Tier 3", "De-prioritise"], 6):
        sm.cell(i, 4, t).fill = tier_fill[t]; sm.cell(i, 5, tier_ct.get(t, 0))
    sm["D12"] = "By Source"; sm["D12"].font = Font(name=A, bold=True, size=12)
    for c, h in zip((4, 5), ["Source", "Contacts"]):
        cc = sm.cell(13, c, h); cc.font = Font(name=A, bold=True, color="FFFFFF"); cc.fill = hdr_fill
    for i, s in enumerate(["Navigator", "Public"], 14):
        sm.cell(i, 4, s); sm.cell(i, 5, src_ct.get(s, 0))
    for col, w in zip("ABCDE", [24, 12, 4, 22, 12]):
        sm.column_dimensions[col].width = w

    dm = sorted([r for r in rows if r["level"] >= 4], key=lambda r: (r["company"], -r["level"], r["title"]))
    tbl(wb.create_sheet("Decision-Makers"), dm)
    tbl(wb.create_sheet("All Contacts"), sorted(rows, key=lambda r: (trank(r["tier"]), r["company"], r["name"])))
    for acc in accounts:
        subset = sorted([r for r in rows if r["company"] == acc], key=lambda r: (-r["level"], r["sr"], r["name"]))
        tab = acc.replace("/", "-").replace("Bristol Myers Squibb", "BMS")[:31]
        tbl(wb.create_sheet(tab), subset)
    wb.save(DATA / "contacts_all.xlsx")
    print("decision-makers tab:", len(dm), "| tabs:", wb.sheetnames)


if __name__ == "__main__":
    main()
